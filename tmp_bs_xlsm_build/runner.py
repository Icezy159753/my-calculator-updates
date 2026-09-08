import contextlib
import importlib.util
import io
import json
import os
import sys
import traceback

import openpyxl
import pandas as pd


SETTING_COLUMNS = (
    'C', 'A', 'S', 'P', 'E', 'AgreeS', 'AgreeP', 'Filter_Var')


def required_columns_from_frame(frame):
    required = set()
    for column in SETTING_COLUMNS:
        if column not in frame.columns:
            continue
        for value in frame[column].dropna():
            text = str(value).strip()
            if not text:
                continue
            for item in text.split(','):
                item = item.strip()
                if item:
                    required.add(item)
    return required


def setting_required_columns(setting_path):
    frame = pd.read_excel(setting_path, sheet_name='Settings')
    return required_columns_from_frame(frame)


def select_setting(settings_dir, raw_columns):
    candidates = []
    for filename in sorted(os.listdir(settings_dir)):
        if not filename.lower().endswith('.xlsx'):
            continue
        path = os.path.join(settings_dir, filename)
        required = setting_required_columns(path)
        missing = sorted(required.difference(raw_columns))
        candidates.append((len(missing), -len(required), filename, path, missing))
    if not candidates:
        raise RuntimeError('No embedded Setting workbook was found')
    candidates.sort()
    missing_count, _, filename, path, missing = candidates[0]
    if missing_count:
        preview = ', '.join(missing[:20])
        raise RuntimeError(
            'Rawdata does not match Setting 1, 4, 5, or 6. '
            'Closest setting is ' + filename + '; missing: ' + preview)
    return path, filename


def append_frame(worksheet, frame):
    worksheet.append(list(frame.columns))
    for row in frame.itertuples(index=False, name=None):
        worksheet.append([
            None if pd.isna(value) else value
            for value in row
        ])


def select_workbook_setting(workbook_path, raw_columns, workdir):
    settings = pd.read_excel(
        workbook_path, sheet_name='Setting Variables')
    labels = pd.read_excel(
        workbook_path, sheet_name='Setting Labels')
    if 'Set_ID' not in settings.columns:
        raise RuntimeError('Setting Variables has no Set_ID column')
    if 'Set_ID' not in labels.columns:
        raise RuntimeError('Setting Labels has no Set_ID column')

    candidates = []
    for set_id, group in settings.groupby('Set_ID', sort=True):
        active = group.drop(columns=['Set_ID']).reset_index(drop=True)
        required = required_columns_from_frame(active)
        missing = sorted(required.difference(raw_columns))
        display_id = str(int(set_id)) if float(set_id).is_integer() else str(set_id)
        candidates.append((
            len(missing), -len(required), display_id, active, missing))
    if not candidates:
        raise RuntimeError('Setting Variables contains no Setting rows')
    candidates.sort(key=lambda item: (item[0], item[1], item[2]))
    missing_count, _, display_id, active, missing = candidates[0]
    if missing_count:
        preview = ', '.join(missing[:20])
        raise RuntimeError(
            'Rawdata does not match Setting 1, 4, 5, or 6. '
            'Closest workbook Set is ' + display_id + '; missing: ' + preview)

    label_group = labels[
        labels['Set_ID'].astype(str).str.replace('.0', '', regex=False)
        == display_id
    ].drop(columns=['Set_ID']).reset_index(drop=True)
    active_path = os.path.join(workdir, 'active_setting.xlsx')
    workbook = openpyxl.Workbook()
    settings_sheet = workbook.active
    settings_sheet.title = 'Settings'
    append_frame(settings_sheet, active)
    labels_sheet = workbook.create_sheet('Label')
    append_frame(labels_sheet, label_group)
    workbook.save(active_path)
    return active_path, 'Set ' + display_id + ' (workbook)'


def patch_setting_path(setting_path, sav_path):
    workbook = openpyxl.load_workbook(setting_path)
    worksheet = workbook['Settings']
    headers = {
        str(cell.value).strip(): cell.column
        for cell in worksheet[1]
        if cell.value is not None
    }
    column = headers.get('PathFile')
    if column is None:
        raise RuntimeError('Settings sheet has no PathFile column')
    worksheet.cell(row=2, column=column).value = sav_path
    workbook.save(setting_path)


def load_module(program_path):
    spec = importlib.util.spec_from_file_location(
        'brandsense_embedded_engine', program_path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


if __name__ == '__main__':
    try:
        # Parse without clever unpacking so Windows paths remain untouched.
        program_path = os.path.abspath(sys.argv[1])
        settings_source = os.path.abspath(sys.argv[2])
        sav_path = os.path.abspath(sys.argv[3])
        raw_csv = os.path.abspath(sys.argv[4])
        mode = sys.argv[5]
        workdir = os.path.abspath(sys.argv[6])
        os.makedirs(workdir, exist_ok=True)
        os.environ['QT_QPA_PLATFORM'] = 'offscreen'

        rawdata = pd.read_csv(raw_csv, low_memory=False)
        if rawdata.empty:
            raise RuntimeError('Rawdata sheet contains no data rows')
        if os.path.isdir(settings_source):
            setting_path, setting_name = select_setting(
                settings_source, set(rawdata.columns))
        else:
            setting_path, setting_name = select_workbook_setting(
                settings_source, set(rawdata.columns), workdir)
        patch_setting_path(setting_path, sav_path)

        module = load_module(program_path)
        from PyQt6.QtWidgets import QApplication
        qapp = QApplication.instance() or QApplication([])
        window = module.SpssProcessorApp()

        window._load_settings_file(
            setting_path, require_pathfile=True)
        window.load_spss_file(
            filepath=sav_path, raise_on_error=True)
        window._snapshot_ui_inputs()

        # Metadata SAV can contain columns not used by Set4. Only require
        # columns referenced by settings or respondent/filter identifiers.
        required = set(window.c_vars_to_compute)
        for values in window.vars_to_transform.values():
            required.update(values)
        required.update(window.id_vars)
        required.update(window._cross_filters())
        missing_required = sorted(
            column for column in required
            if column and column not in rawdata.columns)
        if missing_required:
            preview = ', '.join(missing_required[:20])
            raise RuntimeError(
                'Rawdata is missing required columns: ' + preview)

        window.df = rawdata.copy()
        window.original_filepath = sav_path
        window._snapshot_ui_inputs()
        window._transform_pipeline(with_compute_c=True)
        cross_filters = window._cross_filters()

        qc_candidates = window._build_long_qc_candidates()
        window.qc_candidates_df = qc_candidates
        if mode == 'safe_all':
            window._apply_qc_candidate_rows(
                qc_candidates.index.tolist())
            window._qc_use_model_safeguards = True
        elif mode == 'normal':
            window._restore_qc_full_long_data()
            window._qc_use_model_safeguards = False
        else:
            raise RuntimeError('Unknown run mode: ' + mode)

        output_stub = os.path.join(workdir, 'BrandSense_Engine_Input.sav')
        window.original_filepath = output_stub
        captured = io.StringIO()
        with contextlib.redirect_stdout(captured):
            window._analysis_pipeline(cross_filters)

        output_path = window.last_excel_filepath
        if not output_path or not os.path.exists(output_path):
            raise RuntimeError('Engine did not create an Excel output')

        result = {
            'ok': True,
            'output': output_path,
            'mode': mode,
            'setting': setting_name,
            'raw_rows': int(len(rawdata)),
            'qc_candidates': int(len(qc_candidates)),
            'qc_excluded': int(len(window.qc_excluded_df)),
            'model_warnings': int(len(window._model_warning_groups)),
        }
        with open(
                os.path.join(workdir, 'result.json'),
                'w', encoding='utf-8') as handle:
            json.dump(result, handle, ensure_ascii=False, indent=2)
        with open(
                os.path.join(workdir, 'engine.log'),
                'w', encoding='utf-8') as handle:
            handle.write(captured.getvalue())
        window.close()
        sys.exit(0)
    except Exception as exc:
        workdir = (
            os.path.abspath(sys.argv[6])
            if len(sys.argv) > 6 else os.getcwd())
        os.makedirs(workdir, exist_ok=True)
        with open(
                os.path.join(workdir, 'error.txt'),
                'w', encoding='utf-8') as handle:
            handle.write(f'{type(exc).__name__}: {exc}\n\n')
            handle.write(traceback.format_exc())
        sys.exit(1)
