import tkinter as tk
from tkinter import filedialog, messagebox
import openpyxl
import re
from collections import defaultdict
from openpyxl.styles import Font
import sys

# ===================== Utilities =====================

def clean_header(text):
    if not isinstance(text, str): return ""
    return "".join(re.findall(r"[A-Za-z]", text)).upper()

def parse_sig_groups(sig_input):
    # รองรับ "ABCDEF,GHIJKL,..." และช่วง "A-F,G-L,..."
    groups = [g.strip().upper() for g in sig_input.split(",") if g.strip()]
    expanded = []
    for g in groups:
        if "-" in g and len(g) == 3:
            a,b = g.split("-")
            expanded.append("".join(chr(c) for c in range(ord(a), ord(b)+1)))
        else:
            expanded.append("".join(sorted(set(re.findall(r"[A-Z]", g)))))
    out = {}
    for rng in expanded:
        for ch in rng:
            out[ch] = rng
    return out

def detect_letter_header_row(ws, start_row, end_row, min_seq=3):
    # หาแถวที่มี A,B,C ต่อกัน
    scan_until = min(end_row, start_row + 120)
    for r in range(start_row, scan_until + 1):
        letters = []
        for c in range(3, ws.max_column + 1):  # เริ่มจากคอลัมน์ C
            t = clean_header(str(ws.cell(row=r, column=c).value))
            if len(t) == 1 and 'A' <= t <= 'Z':
                letters.append((c, t))
        if len(letters) < min_seq: 
            continue
        letters.sort()
        vals = {c:v for c,v in letters}
        for c,v in letters:
            if v == 'A' and vals.get(c+1) == 'B' and vals.get(c+2) == 'C':
                return r, c
    return -1, -1

def build_col_letter_map(ws, header_row, start_col):
    mapping = {}
    for c in range(start_col, ws.max_column + 1):
        t = clean_header(str(ws.cell(row=header_row, column=c).value))
        if len(t) == 1 and 'A' <= t <= 'Z':
            mapping[c] = t
        else:
            if mapping: break
    return mapping

def find_total_row_in_colB(ws, header_row, end_row):
    # หา TOTAL ในคอลัมน์ B เท่านั้น (ตามที่พี่ต้องการ)
    for r in range(header_row, end_row + 1):
        v = ws.cell(row=r, column=2).value  # column B = 2
        if isinstance(v, str) and v.strip().upper() == "TOTAL":
            return r
    return -1

def find_last_data_row(ws, start_row, col_indexes):
    last = start_row
    for r in range(start_row, ws.max_row + 1):
        any_val = False
        for c in col_indexes:
            v = ws.cell(row=r, column=c).value
            if v not in (None, "") and str(v).strip() != "":
                any_val = True
                break
        if any_val: last = r
    return last

def filter_sig_text(text, allowed_letters, preserve_tokens=("ADJ",)):
    """
    ลบเฉพาะตัวอักษรอังกฤษที่ 'ไม่อยู่' ใน allowed_letters
    กัน token ใน preserve_tokens เฉพาะกรณีที่ 'ทุกตัวอักษร' ของ token อยู่ใน allowed_letters ด้วย
    ตัวอย่าง: allowed = {'A','B','C','D','E','F'}  -> 'ADJ' กลายเป็น 'AD'
    """
    if not isinstance(text, str) or not text:
        return text

    out = text

    # กำหนด token ที่จะกันแบบมีเงื่อนไข (subset ของ allowed_letters เท่านั้น)
    dynamic_preserve = tuple(
        tok for tok in preserve_tokens
        if tok and set(tok.upper()).issubset(set(allowed_letters))
    )

    # กัน token ด้วย placeholder ที่ไม่ใช่ตัวอักษร
    placeholders = {}
    for idx, tok in enumerate(dynamic_preserve):
        if tok in out:
            ph = f"§{idx}§"  # ใช้สัญลักษณ์ไม่เกี่ยวกับ A-Z
            placeholders[ph] = tok
            out = out.replace(tok, ph)

    # ลบขีดล่างที่อาจตกค้าง
    out = out.replace("_", "")

    # กรองตัวอักษรอังกฤษตาม allowed_letters (อย่างอื่นคงไว้)
    keep = set(allowed_letters)
    buf = []
    for ch in out:
        if 'A' <= ch.upper() <= 'Z':
            if ch.upper() in keep:
                buf.append(ch)
            # ไม่อยู่ใน allowed -> ตัดทิ้ง
        else:
            buf.append(ch)
    out = "".join(buf)

    # คืน token ที่กันไว้
    for ph, tok in placeholders.items():
        out = out.replace(ph, tok)

    return out.strip()


# ===================== Core Processing =====================


def place_sig_stamp(ws, header_row, sig_text, col=3):
    """
    วางข้อความ Sig Used: ... ที่คอลัมน์ C (col=3)
    ให้ขึ้นไปเหนือหัวตาราง 2 แถว (ตามภาพ) ถ้าน้อยกว่าแถว 1 จะถอยลงมาเป็นหัวตารางเอง
    """
    stamp_row = header_row - 2 if header_row - 2 >= 1 else header_row
    cell = ws.cell(row=stamp_row, column=col)
    cell.value = sig_text
    cell.font = Font(name='Arial', size= 9)
# ====== วางทับฟังก์ชันนี้ทั้งก้อน: ปรับเฉพาะโหมด MATRIX ======
# ---------- วางทับฟังก์ชันนี้ทั้งก้อน ----------
def process_single_excel_file(file_path, sig_input, crosstab_mode="NORMAL"):
    import openpyxl
    try:
        # กลุ่มตัวอักษรอนุญาตต่อคอลัมน์
        rules_by_char = parse_sig_groups(sig_input)
        wb = openpyxl.load_workbook(file_path)
        sheets_to_process = [n for n in wb.sheetnames if n.lower() not in ['contents', 'info']]
        cells_changed_count = 0

        for sheet_name in sheets_to_process:
            ws = wb[sheet_name]

            # หาเริ่มตารางด้วยคำว่า Stub ที่คอลัมน์ A (ไม่มีถือว่าชีตเดียวทั้งหน้า)
            stub_starts = [r for r, cell in enumerate(ws['A'], 1) if str(cell.value).startswith('Stub')]
            if not stub_starts:
                stub_starts = [1]
            table_boundaries = stub_starts + [ws.max_row + 2]

            for i in range(len(table_boundaries) - 1):
                start_row = table_boundaries[i]
                end_row   = table_boundaries[i + 1] - 2
                if end_row <= start_row:
                    continue

                sig_text = f"Sig Used: {sig_input}"

                if crosstab_mode.upper() == "NORMAL":
                    # --- หา header แบบเดิม (D= A, E= B) ---
                    header_row, data_total_row = -1, -1
                    for r in range(start_row, end_row + 1):
                        d = ws.cell(row=r, column=4).value
                        e = ws.cell(row=r, column=5).value
                        c = ws.cell(row=r, column=3).value
                        if clean_header(str(d)) == 'A' and clean_header(str(e)) == 'B':
                            header_row = r
                        if str(c).strip().upper() == 'TOTAL':
                            data_total_row = r
                    if header_row == -1 or data_total_row == -1:
                        continue

                    # --- วางแสตมป์ตามภาพ (เหนือหัวตาราง 2 แถว) ---
                    place_sig_stamp(ws, header_row, sig_text, col=3)

                    # --- ลบ Sig ใต้หัวตาราง ---
                    for col in range(3, ws.max_column + 1):
                        hdr = clean_header(str(ws.cell(row=header_row, column=col).value))
                        if hdr in rules_by_char:
                            keep = set(rules_by_char[hdr])
                            for r in range(data_total_row + 1, end_row + 1):
                                v = ws.cell(row=r, column=col).value
                                if isinstance(v, str) and v.strip():
                                    nv = filter_sig_text(v, keep, preserve_tokens=("ADJ",))
                                    if nv != v:
                                        ws.cell(row=r, column=col).value = nv
                                        cells_changed_count += 1
                    continue

                # ================= MATRIX =================
                # 1) หาแถวหัวคอลัมน์ A,B,C ต่อกัน
                header_row, first_col = detect_letter_header_row(ws, start_row, end_row, min_seq=3)
                if header_row == -1:
                    continue

                # วางแสตมป์ตามภาพ
                place_sig_stamp(ws, header_row, sig_text, col=3)

                # 2) map คอลัมน์ -> ตัวอักษร
                col_map = build_col_letter_map(ws, header_row, first_col)
                if not col_map:
                    continue

                # 3) เริ่มประมวลผลตั้งแต่บรรทัดที่คอลัมน์ B = TOTAL (ตามข้อกำหนดล่าสุด)
                total_row = find_total_row_in_colB(ws, header_row, end_row)
                if total_row == -1:
                    continue
                data_start = total_row
                data_end   = find_last_data_row(ws, data_start, list(col_map.keys()))
                if data_end < data_start:
                    continue

                # 4) ลบเฉพาะตัวอักษรอังกฤษใต้หัวตาราง
                for c, letter in col_map.items():
                    allowed = set(rules_by_char.get(letter, ""))
                    for r in range(data_start, data_end + 1):
                        v = ws.cell(row=r, column=c).value
                        if isinstance(v, str) and v.strip():
                            nv = filter_sig_text(v, allowed, preserve_tokens=("ADJ",))
                            if nv != v:
                                ws.cell(row=r, column=c).value = nv
                                cells_changed_count += 1

        wb.save(file_path)
        return cells_changed_count, "Success"

    except Exception as e:
        return 0, f"Error: {e}"



# ===================== GUI (เพิ่มตัวเลือก Crosstab) =====================

class App:
    def __init__(self, root):
        self.root = root
        self.root.title("โปรแกรมลบ Sig จาก Table Lychee V1")

        # Center window
        ww, wh = 700, 450
        sw, sh = self.root.winfo_screenwidth(), self.root.winfo_screenheight()
        self.root.geometry(f"{ww}x{wh}+{int(sw/2-ww/2)}+{int(sh/2-wh/2)}")
        self.root.minsize(650, 400)

        self.file_paths = []
        self.BG_COLOR = "#F0F0F0"
        self.BTN_COLOR_BLUE = "#0078D7"
        self.BTN_HOVER_BLUE = "#005A9E"
        self.BTN_COLOR_ORANGE = "#00AA44"
        self.BTN_COLOR_RED = "#D32F2F"
        self.TEXT_COLOR = "#202020"
        self.HINT_COLOR = "grey"
        self.FONT_NORMAL = ("Segoe UI", 10)
        self.FONT_BOLD = ("Segoe UI", 11, "bold")
        self.HINT_TEXT = "e.g., ABCDEF,GHIJKL,MNOPQR,STUVWX,YZ"

        self.root.configure(bg=self.BG_COLOR)
        self.root.grid_columnconfigure(0, weight=1)
        self.root.grid_rowconfigure(3, weight=1)

        # Top buttons
        top = tk.Frame(root, bg=self.BG_COLOR); top.grid(row=0, column=0, sticky="ew", padx=10, pady=(10,5))
        tk.Button(top, text="เลือกไฟล์...", font=self.FONT_BOLD, command=self.browse_files,
                  bg=self.BTN_COLOR_ORANGE, fg="white", relief="flat").pack(side=tk.LEFT)
        tk.Button(top, text="ลบไฟล์ที่เลือก", font=self.FONT_BOLD, command=self.delete_selected_files,
                  bg=self.BTN_COLOR_RED, fg="white", relief="flat").pack(side=tk.LEFT, padx=(10,0))

        # Mode
        mode = tk.Frame(root, bg=self.BG_COLOR); mode.grid(row=1, column=0, sticky="w", padx=10, pady=(0,0))
        tk.Label(mode, text="ประเภท Crosstab :", font=self.FONT_NORMAL, bg=self.BG_COLOR).pack(side=tk.LEFT, padx=(0,6))
        self.mode_var = tk.StringVar(value="NORMAL")
        tk.Radiobutton(mode, text="Crosstab ธรรมดา", variable=self.mode_var, value="NORMAL",
                       font=self.FONT_NORMAL, bg=self.BG_COLOR).pack(side=tk.LEFT)
        tk.Radiobutton(mode, text="Matrix", variable=self.mode_var, value="MATRIX",
                       font=self.FONT_NORMAL, bg=self.BG_COLOR).pack(side=tk.LEFT, padx=(10,0))

        # Sig input
        sigf = tk.Frame(root, bg=self.BG_COLOR); sigf.grid(row=2, column=0, sticky="ew", padx=10, pady=5)
        tk.Label(sigf, text="ใส่ Sig ตรงนี้ :", font=self.FONT_NORMAL, bg=self.BG_COLOR).grid(row=0, column=0, padx=(0,5))
        self.sig_entry = tk.Entry(sigf, font=self.FONT_NORMAL, relief="solid", bd=1, width=82, fg=self.HINT_COLOR)
        self.sig_entry.grid(row=0, column=1, sticky="w")
        self.sig_entry.insert(0, self.HINT_TEXT)
        self.sig_entry.bind("<FocusIn>", lambda e: (self.sig_entry.delete(0, tk.END), self.sig_entry.config(fg=self.TEXT_COLOR)) if self.sig_entry.get()==self.HINT_TEXT else None)
        self.sig_entry.bind("<FocusOut>", lambda e: (self.sig_entry.insert(0, self.HINT_TEXT), self.sig_entry.config(fg=self.HINT_COLOR)) if not self.sig_entry.get() else None)

        # File list
        lf = tk.Frame(root, bg=self.BG_COLOR); lf.grid(row=3, column=0, sticky="nsew", padx=10)
        lf.grid_rowconfigure(0, weight=1); lf.grid_columnconfigure(0, weight=1)
        self.file_listbox = tk.Listbox(lf, font=self.FONT_NORMAL, relief="solid", bd=1, selectmode=tk.EXTENDED)
        self.file_listbox.grid(row=0, column=0, sticky="nsew")
        sb = tk.Scrollbar(lf, orient="vertical", command=self.file_listbox.yview); sb.grid(row=0, column=1, sticky="ns")
        self.file_listbox.config(yscrollcommand=sb.set)

        # Process
        tk.Button(root, text="เริ่มต้นลบ Sig.....", font=self.FONT_BOLD, bg=self.BTN_COLOR_BLUE, fg="white",
                  activebackground=self.BTN_HOVER_BLUE, activeforeground="white", relief="flat",
                  cursor="hand2", command=self.process_files).grid(row=4, column=0, sticky="ew", padx=10, pady=(15,0), ipady=8)

        # Status
        self.status_var = tk.StringVar(value="โปรแกรมพร้อมใช้งาน")
        tk.Label(root, textvariable=self.status_var, font=("Segoe UI", 9), relief=tk.SUNKEN, anchor="w", padx=5, fg="#505050").grid(row=5, column=0, sticky="ew")

    def browse_files(self):
        paths = filedialog.askopenfilenames(filetypes=[("Excel Files", "*.xlsx")])
        if paths:
            self.file_paths = list(paths)
            self.file_listbox.delete(0, tk.END)
            for p in self.file_paths:
                self.file_listbox.insert(tk.END, " " + p.split("/")[-1])
            self.status_var.set(f"{len(self.file_paths)} ไฟล์ที่เลือก พร้อมจะลบ Sig.")

    def delete_selected_files(self):
        idxs = self.file_listbox.curselection()
        if not idxs: return
        for i in sorted(idxs, reverse=True):
            self.file_listbox.delete(i)
            del self.file_paths[i]
        self.status_var.set(f"{len(self.file_paths)} files remaining.")

    def process_files(self):
        if not self.file_paths:
            messagebox.showwarning("Warning", "Please select at least one Excel file.")
            return
        sig_groups = self.sig_entry.get()
        if not sig_groups or sig_groups == self.HINT_TEXT:
            messagebox.showwarning("Warning", "Please enter Sig groups.")
            return

        mode = self.mode_var.get().upper()
        total_files = len(self.file_paths)
        processed_files, total_cells_changed, had_errors = 0, 0, False

        for i, file_path in enumerate(self.file_paths):
            fname = file_path.split('/')[-1]
            self.status_var.set(f"Processing ({mode}) file {i+1}/{total_files}: {fname}...")
            self.root.update_idletasks()
            cells_changed, status = process_single_excel_file(file_path, sig_groups, mode)
            if status == "Success":
                processed_files += 1
                total_cells_changed += cells_changed
            else:
                had_errors = True
                messagebox.showerror("Processing Error", f"An error occurred while processing:\n{fname}\n\nDetails: {status}")
                break

        msg = f"ลบ Sig เรียบร้อย [{mode}].\n\nสำเร็จทั้ง: {processed_files}/{total_files} files.\nลบไปทั้งหมด: {total_cells_changed} cells."
        if had_errors:
            msg += "\n\nSome files encountered errors and were not processed."
            messagebox.showwarning("Process Finished with Errors", msg)
        else:
            messagebox.showinfo("Success", msg)

        self.file_listbox.delete(0, tk.END)
        self.file_paths.clear()
        self.sig_entry.delete(0, tk.END)
        self.sig_entry.insert(0, self.HINT_TEXT)
        self.status_var.set("เริ่มต้นลบ Sig.....รอบต่อไป")

# -------- Entry point --------
def run_this_app(working_dir=None):
    try:
        root = tk.Tk()
        app = App(root)
        root.mainloop()
    except Exception as e:
        if 'root' not in locals() or not root.winfo_exists():
            rt = tk.Tk(); rt.withdraw()
            messagebox.showerror("Application Error", f"An unexpected error occurred:\n{e}", parent=rt)
            rt.destroy()
        else:
            messagebox.showerror("Application Error", f"An unexpected error occurred:\n{e}", parent=root)
        sys.exit(f"Error running app: {e}")

if __name__ == "__main__":
    run_this_app()
